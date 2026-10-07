import json
import math
import sys
import unicodedata

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

SCRIPTS_DIR = sys.argv[1]
OUT = sys.argv[2]

FONT_NAME = "微软雅黑"
BODY = 10
TEAL = "0F6E56"

thin = Side(style="thin", color="B4B2A9")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)


def font(bold=False, size=BODY, color="000000"):
    return Font(name=FONT_NAME, size=size, bold=bold, color=color)


F_BODY = font()
F_BOLD = font(bold=True)
F_HEAD = font(bold=True, color="FFFFFF")
F_TITLE = font(bold=True, size=14, color=TEAL)
F_NOTE = font(color="5F5E5A")
F_INPUT = font(color="0000FF")
F_LINK = font(color="008000")
F_LINK_BOLD = font(bold=True, color="008000")

FILL_HEAD = PatternFill("solid", fgColor=TEAL)
FILL_SUB = PatternFill("solid", fgColor="E1F5EE")
FILL_TOTAL = PatternFill("solid", fgColor="F1EFE8")
FILL_INPUT = PatternFill("solid", fgColor="FFF2CC")

WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(wrap_text=True, vertical="center", horizontal="center")
MID = Alignment(wrap_text=True, vertical="center")

GROUP_FMT = '"第"0"组"'

GROUPS = {
    1: ("大巴车上推销土特产", "降噪与语音端点检测"),
    2: ("购物店“包厢式”推销", "远距离、手机放口袋时的识别"),
    3: ("擅自变更行程", "说话人分离"),
    4: ("自费项目费用纠纷", "数字、证号、金额规范化"),
    5: ("兜售物品、索要小费", "热词"),
    6: ("言语施压、强迫消费（措辞温和）", "TensorFlow 话术分类（模型 A）"),
    7: ("服务态度纠纷", "TensorFlow 话术分类（模型 B）"),
    8: ("正常讲解（对照组）", "时间戳与证据片段导出、误报率"),
}
LABELS = [
    ("购物安排", "推荐、安排、推销购物或进店，商品介绍，游客对购物的回应"),
    ("费用", "价格、自费项目、团费、退款、小费等金额相关"),
    ("行程变更", "取消、压缩、调换行程或时间安排变动"),
    ("服务态度", "催促、敷衍、不耐烦，以及对态度的抱怨和道歉"),
    ("威胁消费", "以行程、住宿、服务等为条件含蓄施压要求消费（措辞温和，无辱骂、无人身威胁）"),
    ("正常讲解", "景点介绍、安全提醒、集合通知等正常导游工作"),
    ("其他", "寒暄、问路、附和，以及与上面无关的话"),
]
CONDITIONS = [
    ("安静", "Q", "安静的教室或会议室；手机平放在说话人中间的桌上，距离约 50 厘米"),
    ("嘈杂教室", "N", "课间或几组同时录音的教室；手机位置同上"),
    ("口袋或远距离", "F", "手机放在一名演员的外衣口袋里，或放在离说话人 2—3 米外的桌角"),
]
NOTE_G6S3 = "游客复述导游原话：按说话人意图标为“服务态度”；如按话题标注，可改为“威胁消费”"
G6S3_QUOTED = {11, 12, 13, 14, 27}

PUNCT = ["，", "。", "？", "！", "、", "…", "；", "：", "《", "》", " ", "“", "”", "—", "（", "）", "·"]
PUNCT_ARRAY = "{" + ",".join(f'"{p}"' for p in PUNCT) + "}"


def disp_len(text):
    if text is None:
        return 0
    s = str(text)
    if s.startswith("="):
        return 6
    return sum(2 if unicodedata.east_asian_width(ch) in ("W", "F") else 1 for ch in s)


def cell(ws, r, c, v, align=WRAP, fnt=None, fill=None, fmt=None):
    x = ws.cell(row=r, column=c, value=v)
    x.font = fnt or F_BODY
    x.alignment = align
    x.border = BORDER
    if fill:
        x.fill = fill
    if fmt:
        x.number_format = fmt
    return x


def header(ws, row, labels):
    for i, lab in enumerate(labels, start=1):
        c = ws.cell(row=row, column=i, value=lab)
        c.font = F_HEAD
        c.fill = FILL_HEAD
        c.alignment = CENTER
        c.border = BORDER
    ws.row_dimensions[row].height = 32


def title(ws, text, note=None, span=8, note_width=None):
    ws["A1"] = text
    ws["A1"].font = F_TITLE
    ws.row_dimensions[1].height = 26
    if note:
        ws["A2"] = note
        ws["A2"].font = F_NOTE
        ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=span)
        if note_width:
            n = sum(max(1, math.ceil(disp_len(p) / max((note_width - 2) * 0.9, 2))) for p in note.split("\n"))
            ws.row_dimensions[2].height = n * 15 + 6


def widths(ws, ws_widths):
    for i, w in enumerate(ws_widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def fit_rows(ws, first, last, line_pt=15, pad=6, skip_cols=()):
    merged = set()
    for rng in ws.merged_cells.ranges:
        if rng.min_row != rng.max_row or rng.min_col != rng.max_col:
            for r in range(rng.min_row, rng.max_row + 1):
                for c in range(rng.min_col, rng.max_col + 1):
                    merged.add((r, c))
    for r in range(first, last + 1):
        lines = 1
        for x in ws[r]:
            if x.value is None or (r, x.column) in merged or x.column in skip_cols:
                continue
            w = ws.column_dimensions[get_column_letter(x.column)].width or 8.43
            usable = max((w - 1.5) * 0.9, 2)
            n = sum(max(1, math.ceil(disp_len(p) / usable)) for p in str(x.value).split("\n"))
            lines = max(lines, n)
        ws.row_dimensions[r].height = lines * line_pt + pad


# ------------------------------------------------------------------ load scripts
scripts = []
for g in range(1, 9):
    with open(f"{SCRIPTS_DIR}/group_{g}.json", encoding="utf-8") as f:
        data = json.load(f)
    for sc in data["scripts"]:
        scripts.append((g, sc))
n_lines = sum(len(sc["lines"]) for _, sc in scripts)

wb = Workbook()
ws_i = wb.active
ws_i.title = "说明"
ws_o = wb.create_sheet("剧本总览")
ws_l = wb.create_sheet("台词")
ws_g = wb.create_sheet("分组统计")
ws_r = wb.create_sheet("录音计划")
ws_n = wb.create_sheet("虚构名称")

# ------------------------------------------------------------------ 台词
L_FIRST = 5
L_LAST = L_FIRST + n_lines - 1
title(
    ws_l,
    "全部台词（录音参考文本）",
    "每句一行。“动作与环境提示”一栏不念，只用于表演。允许即兴改口，但录完必须对照录音修改“台词”列，"
    "以实际说了什么为准；修改后有效字数和预计时长会自动重算。可用表头的筛选按钮按剧本、组或话术类别筛选。",
    span=11,
    note_width=210,
)
header(ws_l, 4, ["剧本编号", "组", "行号", "说话人", "台词（参考文本）", "动作与环境提示（不念）",
                 "话术类别", "数字规范写法", "热词", "有效字数", "标注备注"])
r = L_FIRST
for g, sc in scripts:
    for i, ln in enumerate(sc["lines"], start=1):
        cell(ws_l, r, 1, sc["id"], CENTER)
        cell(ws_l, r, 2, g, CENTER, fmt=GROUP_FMT)
        cell(ws_l, r, 3, i, CENTER)
        cell(ws_l, r, 4, ln["speaker"], CENTER)
        cell(ws_l, r, 5, ln["text"])
        cell(ws_l, r, 6, ln["direction"] or None, fnt=F_NOTE)
        cell(ws_l, r, 7, ln["label"], CENTER)
        cell(ws_l, r, 8, ln["numbers"] or None)
        cell(ws_l, r, 9, ln["hotwords"] or None)
        cell(ws_l, r, 10, f'=LEN(E{r})-SUMPRODUCT(LEN(E{r})-LEN(SUBSTITUTE(E{r},{PUNCT_ARRAY},"")))', CENTER)
        note = NOTE_G6S3 if (sc["id"] == "G6-S3" and i in G6S3_QUOTED) else None
        cell(ws_l, r, 11, note, fnt=F_NOTE)
        r += 1
widths(ws_l, [10, 8, 6, 9, 58, 26, 10, 22, 26, 8, 24])
fit_rows(ws_l, L_FIRST, L_LAST, skip_cols=(10,))
ws_l.freeze_panes = "E5"
ws_l.auto_filter.ref = f"A4:K{L_LAST}"

# ------------------------------------------------------------------ 剧本总览
O_FIRST = 8
O_LAST = O_FIRST + len(scripts) - 1
title(
    ws_o,
    "剧本总览（8 组 × 3 个剧本）",
    "台词行数、有效字数、预计时长由公式根据“台词”表自动计算。蓝色数字可改：朗读语速按实际试读情况调整，时长检查会随之更新。",
    span=13,
    note_width=200,
)
o_params = [(3, "朗读语速（字/分钟）", 220), (4, "每个剧本时长下限（分钟）", 3), (5, "每个剧本时长上限（分钟）", 5)]
for rr, lab, val in o_params:
    cell(ws_o, rr, 1, lab, MID, fnt=F_BOLD, fill=FILL_SUB)
    ws_o.merge_cells(start_row=rr, start_column=1, end_row=rr, end_column=3)
    for c in (2, 3):
        ws_o.cell(row=rr, column=c).border = BORDER
    cell(ws_o, rr, 4, val, CENTER, fnt=F_INPUT)
    ws_o.row_dimensions[rr].height = 20
header(ws_o, 7, ["剧本编号", "组", "场景", "标题", "地点与背景", "角色", "本剧本侧重",
                 "台词行数", "有效字数", "预计时长（分钟）", "时长检查", "角色分配（谁演谁）", "状态"])
LR = f"$A${L_FIRST}:$A${L_LAST}"
LJ = f"$J${L_FIRST}:$J${L_LAST}"
for k, (g, sc) in enumerate(scripts):
    r = O_FIRST + k
    cell(ws_o, r, 1, sc["id"], CENTER, fnt=F_BOLD)
    cell(ws_o, r, 2, g, CENTER, fmt=GROUP_FMT)
    cell(ws_o, r, 3, GROUPS[g][0])
    cell(ws_o, r, 4, sc["title"], fnt=F_BOLD)
    cell(ws_o, r, 5, sc["setting"])
    cell(ws_o, r, 6, "、".join(sc["roles"]))
    cell(ws_o, r, 7, sc["focus"])
    cell(ws_o, r, 8, f"=COUNTIF('台词'!{LR},A{r})", CENTER, fnt=F_LINK)
    cell(ws_o, r, 9, f"=SUMIF('台词'!{LR},A{r},'台词'!{LJ})", CENTER, fnt=F_LINK)
    cell(ws_o, r, 10, f"=I{r}/$D$3", CENTER, fmt="0.0")
    cell(ws_o, r, 11, f'=IF(AND(J{r}>=$D$4,J{r}<=$D$5),"符合","需调整")', CENTER)
    cell(ws_o, r, 12, None, fill=FILL_INPUT)
    cell(ws_o, r, 13, "未开始", CENTER, fill=FILL_INPUT)
O_TOTAL = O_LAST + 1
cell(ws_o, O_TOTAL, 1, "合计", CENTER, fnt=F_BOLD, fill=FILL_TOTAL)
for c in range(2, 14):
    cell(ws_o, O_TOTAL, c, None, fill=FILL_TOTAL)
ws_o.cell(row=O_TOTAL, column=8, value=f"=SUM(H{O_FIRST}:H{O_LAST})").font = F_BOLD
ws_o.cell(row=O_TOTAL, column=9, value=f"=SUM(I{O_FIRST}:I{O_LAST})").font = F_BOLD
ws_o.cell(row=O_TOTAL, column=10, value=f"=SUM(J{O_FIRST}:J{O_LAST})").font = F_BOLD
ws_o.cell(row=O_TOTAL, column=10).number_format = "0.0"
ws_o.cell(row=O_TOTAL, column=11, value=f'=COUNTIF(K{O_FIRST}:K{O_LAST},"符合")&" / "&COUNTA(A{O_FIRST}:A{O_LAST})&" 个符合"').font = F_BOLD
for c in (8, 9, 10, 11):
    ws_o.cell(row=O_TOTAL, column=c).alignment = CENTER
widths(ws_o, [9, 8, 18, 16, 36, 22, 40, 8, 8, 9, 8, 26, 10])
fit_rows(ws_o, O_FIRST, O_TOTAL)
ws_o.freeze_panes = "B8"
dv_o = DataValidation(type="list", formula1='"未开始,已分配角色,已录音,已校对"', allow_blank=True)
ws_o.add_data_validation(dv_o)
dv_o.add(f"M{O_FIRST}:M{O_LAST}")

# ------------------------------------------------------------------ 分组统计
title(
    ws_g,
    "分组统计",
    "检查每组每种录音条件是否有 10 分钟以上、三种条件合计是否在 30—45 分钟，以及全班合计是否在 4—6 小时。蓝色数字为可调整的目标值。",
    span=10,
    note_width=150,
)
g_params = [
    (3, "每种条件最少（分钟）", 10),
    (4, "三种条件合计下限（分钟）", 30),
    (5, "三种条件合计上限（分钟）", 45),
    (6, "录音条件数", 3),
    (7, "全班合计下限（小时）", 4),
    (8, "全班合计上限（小时）", 6),
]
for rr, lab, val in g_params:
    cell(ws_g, rr, 1, lab, MID, fnt=F_BOLD, fill=FILL_SUB)
    cell(ws_g, rr, 2, val, CENTER, fnt=F_INPUT)
    ws_g.row_dimensions[rr].height = 20

G_HEAD = 10
header(ws_g, G_HEAD, ["组", "场景", "专项", "剧本数", "台词行数", "有效字数",
                      "每种条件预计时长（分钟）", "每种条件是否够 10 分钟", "三种条件合计（分钟）", "合计是否在 30—45 分钟"])
G_FIRST = G_HEAD + 1
OB = f"'剧本总览'!$B${O_FIRST}:$B${O_LAST}"
for k, g in enumerate(range(1, 9)):
    r = G_FIRST + k
    cell(ws_g, r, 1, g, CENTER, fnt=F_BOLD, fmt=GROUP_FMT)
    cell(ws_g, r, 2, GROUPS[g][0])
    cell(ws_g, r, 3, GROUPS[g][1])
    cell(ws_g, r, 4, f"=COUNTIF({OB},A{r})", CENTER, fnt=F_LINK)
    cell(ws_g, r, 5, f"=SUMIF({OB},A{r},'剧本总览'!$H${O_FIRST}:$H${O_LAST})", CENTER, fnt=F_LINK)
    cell(ws_g, r, 6, f"=SUMIF({OB},A{r},'剧本总览'!$I${O_FIRST}:$I${O_LAST})", CENTER, fnt=F_LINK)
    cell(ws_g, r, 7, f"=SUMIF({OB},A{r},'剧本总览'!$J${O_FIRST}:$J${O_LAST})", CENTER, fnt=F_LINK, fmt="0.0")
    cell(ws_g, r, 8, f'=IF(G{r}>=$B$3,"够","不够")', CENTER)
    cell(ws_g, r, 9, f"=G{r}*$B$6", CENTER, fmt="0.0")
    cell(ws_g, r, 10, f'=IF(AND(I{r}>=$B$4,I{r}<=$B$5),"符合","需调整")', CENTER)
G_LAST = G_FIRST + 7
G_TOTAL = G_LAST + 1
cell(ws_g, G_TOTAL, 1, "全班合计", CENTER, fnt=F_BOLD, fill=FILL_TOTAL)
for c in range(2, 11):
    cell(ws_g, G_TOTAL, c, None, CENTER, fill=FILL_TOTAL)
for c, col in ((4, "D"), (5, "E"), (6, "F"), (7, "G"), (9, "I")):
    x = ws_g.cell(row=G_TOTAL, column=c, value=f"=SUM({col}{G_FIRST}:{col}{G_LAST})")
    x.font = F_BOLD
    if c in (7, 9):
        x.number_format = "0.0"
G_HOURS = G_TOTAL + 1
cell(ws_g, G_HOURS, 1, "全班合计（小时）", CENTER, fnt=F_BOLD, fill=FILL_TOTAL)
for c in range(2, 11):
    cell(ws_g, G_HOURS, c, None, CENTER, fill=FILL_TOTAL)
x = ws_g.cell(row=G_HOURS, column=9, value=f"=I{G_TOTAL}/60")
x.font = F_BOLD
x.number_format = "0.0"
ws_g.cell(row=G_HOURS, column=10, value=f'=IF(AND(I{G_HOURS}>=$B$7,I{G_HOURS}<=$B$8),"符合","需调整")').font = F_BOLD

# label distribution
D_HEAD = G_HOURS + 3
ws_g.cell(row=D_HEAD - 1, column=1, value="话术类别分布（台词行数）").font = F_BOLD
hdr =ws_g.cell(row=D_HEAD, column=1, value="话术类别")
hdr.font, hdr.fill, hdr.alignment, hdr.border = F_HEAD, FILL_HEAD, CENTER, BORDER
for k, g in enumerate(range(1, 9)):
    h = ws_g.cell(row=D_HEAD, column=2 + k, value=g)
    h.font, h.fill, h.alignment, h.border = F_HEAD, FILL_HEAD, CENTER, BORDER
    h.number_format = GROUP_FMT
h = ws_g.cell(row=D_HEAD, column=10, value="合计")
h.font, h.fill, h.alignment, h.border = F_HEAD, FILL_HEAD, CENTER, BORDER
ws_g.row_dimensions[D_HEAD].height = 24
LB = f"'台词'!$B${L_FIRST}:$B${L_LAST}"
LG = f"'台词'!$G${L_FIRST}:$G${L_LAST}"
for k, (lab, _) in enumerate(LABELS):
    r = D_HEAD + 1 + k
    cell(ws_g, r, 1, lab, CENTER, fnt=F_BOLD)
    for j in range(8):
        col = get_column_letter(2 + j)
        cell(ws_g, r, 2 + j, f"=COUNTIFS({LB},{col}${D_HEAD},{LG},$A{r})", CENTER, fnt=F_LINK)
    cell(ws_g, r, 10, f"=SUM(B{r}:I{r})", CENTER, fnt=F_BOLD)
D_LAST = D_HEAD + len(LABELS)
D_TOTAL = D_LAST + 1
cell(ws_g, D_TOTAL, 1, "合计", CENTER, fnt=F_BOLD, fill=FILL_TOTAL)
for j in range(9):
    col = get_column_letter(2 + j)
    cell(ws_g, D_TOTAL, 2 + j, f"=SUM({col}{D_HEAD + 1}:{col}{D_LAST})", CENTER, fnt=F_BOLD, fill=FILL_TOTAL)
widths(ws_g, [22, 26, 28, 9, 9, 9, 13, 12, 13, 13])
fit_rows(ws_g, G_FIRST, G_HOURS)
for r in range(D_HEAD + 1, D_TOTAL + 1):
    ws_g.row_dimensions[r].height = 20

# ------------------------------------------------------------------ 录音计划
title(
    ws_r,
    "录音计划（24 个剧本 × 3 种条件 = 72 个录音文件）",
    "文件名 = 剧本编号-条件代码.wav（Q 安静，N 嘈杂教室，F 口袋或远距离）。统一用同一个录音 App，16kHz 单声道 WAV。"
    "录音人、日期、实际时长和两个勾选项为待填写项。",
    span=12,
    note_width=190,
)
header(ws_r, 4, ["录音文件名", "剧本编号", "组", "录音条件", "条件代码", "怎么录", "预计时长（分钟）",
                 "录音人", "录音日期", "实际时长（分钟）", "参考文本已校对", "已放入数据池"])
R_FIRST = 5
r = R_FIRST
OA = f"'剧本总览'!$A${O_FIRST}:$A${O_LAST}"
OJ = f"'剧本总览'!$J${O_FIRST}:$J${O_LAST}"
for g, sc in scripts:
    for cname, code, how in CONDITIONS:
        cell(ws_r, r, 1, f'=B{r}&"-"&E{r}&".wav"', CENTER, fnt=F_BOLD)
        cell(ws_r, r, 2, sc["id"], CENTER)
        cell(ws_r, r, 3, g, CENTER, fmt=GROUP_FMT)
        cell(ws_r, r, 4, cname, CENTER)
        cell(ws_r, r, 5, code, CENTER)
        cell(ws_r, r, 6, how, fnt=F_NOTE)
        cell(ws_r, r, 7, f"=INDEX({OJ},MATCH(B{r},{OA},0))", CENTER, fnt=F_LINK, fmt="0.0")
        for c in (8, 9, 10):
            cell(ws_r, r, c, None, CENTER, fill=FILL_INPUT, fmt="0.0" if c == 10 else None)
        cell(ws_r, r, 11, "否", CENTER, fill=FILL_INPUT)
        cell(ws_r, r, 12, "否", CENTER, fill=FILL_INPUT)
        r += 1
R_LAST = r - 1
dv_r = DataValidation(type="list", formula1='"是,否"', allow_blank=True)
ws_r.add_data_validation(dv_r)
dv_r.add(f"K{R_FIRST}:L{R_LAST}")

S_HEAD = R_LAST + 2
ws_r.cell(row=S_HEAD, column=1, value="汇总").font = F_BOLD
summary = [("安静", f'=SUMIF($D${R_FIRST}:$D${R_LAST},"安静",$G${R_FIRST}:$G${R_LAST})'),
           ("嘈杂教室", f'=SUMIF($D${R_FIRST}:$D${R_LAST},"嘈杂教室",$G${R_FIRST}:$G${R_LAST})'),
           ("口袋或远距离", f'=SUMIF($D${R_FIRST}:$D${R_LAST},"口袋或远距离",$G${R_FIRST}:$G${R_LAST})')]
hrow = S_HEAD + 1
for c, lab in enumerate(["项目", "预计（分钟）", "实际（分钟）", "已完成"], start=1):
    h = ws_r.cell(row=hrow, column=c, value=lab)
    h.font, h.fill, h.alignment, h.border = F_HEAD, FILL_HEAD, CENTER, BORDER
for k, (cname, f_planned) in enumerate(summary):
    rr = hrow + 1 + k
    cell(ws_r, rr, 1, cname, CENTER, fnt=F_BOLD)
    cell(ws_r, rr, 2, f_planned, CENTER, fmt="0.0")
    cell(ws_r, rr, 3, f'=SUMIF($D${R_FIRST}:$D${R_LAST},A{rr},$J${R_FIRST}:$J${R_LAST})', CENTER, fmt="0.0")
    cell(ws_r, rr, 4, f'=COUNTIFS($D${R_FIRST}:$D${R_LAST},A{rr},$L${R_FIRST}:$L${R_LAST},"是")&" / "&COUNTIF($D${R_FIRST}:$D${R_LAST},A{rr})', CENTER)
tr = hrow + 4
cell(ws_r, tr, 1, "合计（分钟）", CENTER, fnt=F_BOLD, fill=FILL_TOTAL)
cell(ws_r, tr, 2, f"=SUM(B{hrow + 1}:B{hrow + 3})", CENTER, fnt=F_BOLD, fill=FILL_TOTAL, fmt="0.0")
cell(ws_r, tr, 3, f"=SUM(C{hrow + 1}:C{hrow + 3})", CENTER, fnt=F_BOLD, fill=FILL_TOTAL, fmt="0.0")
cell(ws_r, tr, 4, f'=COUNTIF($L${R_FIRST}:$L${R_LAST},"是")&" / "&COUNTA($B${R_FIRST}:$B${R_LAST})', CENTER, fnt=F_BOLD, fill=FILL_TOTAL)
hr = tr + 1
cell(ws_r, hr, 1, "合计（小时）", CENTER, fnt=F_BOLD, fill=FILL_TOTAL)
cell(ws_r, hr, 2, f"=B{tr}/60", CENTER, fnt=F_BOLD, fill=FILL_TOTAL, fmt="0.0")
cell(ws_r, hr, 3, f"=C{tr}/60", CENTER, fnt=F_BOLD, fill=FILL_TOTAL, fmt="0.0")
cell(ws_r, hr, 4, None, fill=FILL_TOTAL)
widths(ws_r, [16, 10, 8, 13, 8, 40, 10, 10, 11, 10, 10, 10])
fit_rows(ws_r, R_FIRST, R_LAST)
ws_r.freeze_panes = "B5"
ws_r.auto_filter.ref = f"A4:L{R_LAST}"

# ------------------------------------------------------------------ 虚构名称
title(
    ws_n,
    "虚构名称与号码",
    "剧本中所有旅行社、店铺、酒店、人名、号码均为虚构。如发现与真实企业或人员重名，请统一替换后再使用。"
    "这张表也可直接作为第 5 组热词表的起点。",
    span=4,
    note_width=110,
)
header(ws_n, 4, ["类型", "名称", "出现的组", "说明"])
names = [
    ("旅行社", "云栖野渡旅行社", "第 1 组", "虚构"),
    ("旅行社", "鹿鸣七彩旅行社", "第 2 组", "虚构"),
    ("旅行社", "青岫远途旅行社", "第 3 组", "虚构"),
    ("旅行社", "拾光滇程旅行社", "第 4 组", "虚构"),
    ("旅行社", "雾隐行舟旅行社", "第 5 组", "虚构；剧本中故意出现“雾隐晚渡”“雾隐那家”等说错或简称"),
    ("旅行社", "松风晚渡旅行社", "第 5 组", "虚构；剧本中故意出现“松风行舟”等说错"),
    ("旅行社", "听澜百川旅行社", "第 6 组", "虚构"),
    ("旅行社", "半山观澜旅行社", "第 7 组", "虚构"),
    ("旅行社", "星野云途旅行社", "第 8 组", "虚构"),
    ("店铺", "暮山特产超市", "第 1 组", "虚构"),
    ("店铺", "雾林茶舍", "第 1 组", "虚构"),
    ("店铺", "青黛玉石馆", "第 2 组", "虚构"),
    ("店铺", "竹溪茶仓", "第 2 组", "虚构"),
    ("店铺", "晓月银坊", "第 2、5 组", "虚构；第 5 组剧本中故意出现“晓月阁”等说错"),
    ("店铺", "听松阁玉器行", "第 5 组", "虚构；第 5 组剧本中故意出现“听松坊”等说错"),
    ("店铺", "翠微坊珠宝店", "第 5 组", "虚构"),
    ("店铺", "锦程玉器城", "第 6 组", "虚构"),
    ("品牌", "百年茶语", "第 5 组", "虚构；剧本中故意出现“百年茶叶”的误听"),
    ("酒店", "澄溪假日酒店", "第 3 组", "虚构"),
    ("酒店", "远山客栈", "第 3 组", "虚构"),
    ("酒店", "临湖酒店", "第 7 组", "虚构"),
    ("演出与设施", "《云上花间》", "第 4 组", "虚构演出"),
    ("演出与设施", "望湖索道", "第 4 组", "虚构"),
    ("演出与设施", "环湖电瓶车", "第 4 组", "虚构"),
    ("人物称呼", "周导、马导、李导、陈导、杨导、赵导、刘导、何导", "第 1—8 组各一位", "虚构，只用姓氏称呼"),
    ("人物称呼", "张经理", "第 3 组", "虚构"),
    ("号码", "旅行社电话 0871-0000-6688", "第 4 组", "虚构，含“0000”号段，不对应真实号码"),
    ("号码", "导游证号 YN-0000-3721", "第 4 组", "虚构"),
    ("号码", "合同编号 HT-20260000-118", "第 4 组", "虚构"),
    ("号码", "订单号 DD-0000-5566", "第 4 组", "虚构"),
    ("公共热线", "12301、12345", "多个组", "真实的公共服务热线，剧本中只作为投诉渠道提到"),
]
for k, row in enumerate(names):
    rr = 5 + k
    for c, v in enumerate(row, start=1):
        cell(ws_n, rr, c, v, CENTER if c in (1, 3) else WRAP, fnt=F_BOLD if c == 2 else None)
widths(ws_n, [12, 34, 16, 48])
fit_rows(ws_n, 5, 4 + len(names))

# ------------------------------------------------------------------ 说明
ws_i["A1"] = "测评录音剧本 · 旅游纠纷录音材料整理项目"
ws_i["A1"].font = F_TITLE
ws_i.row_dimensions[1].height = 26
meta = [
    ("整理日期", "2026-10-07"),
    ("内容", f"8 组 × 3 个剧本，共 24 个虚构情景剧，{n_lines} 行台词"),
    ("用途", "照剧本录音，剧本即录音的参考答案；同时提供说话人、话术类别、数字规范写法、热词等标注"),
    ("录音量", "每个剧本 3 种条件各录一遍，共 72 个文件；预计每组每种条件 12—14 分钟，全班合计约 5 小时（见“分组统计”）"),
]
r = 3
for k, v in meta:
    cell(ws_i, r, 1, k, MID, fnt=F_BOLD, fill=FILL_SUB)
    cell(ws_i, r, 2, v)
    ws_i.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
    for c in (3, 4):
        ws_i.cell(row=r, column=c).border = BORDER
    r += 1

r += 1
header_row = r
for c, lab in enumerate(["录音条件", "代码", "怎么录", ""], start=1):
    h = ws_i.cell(row=r, column=c, value=lab)
    h.font, h.fill, h.alignment, h.border = F_HEAD, FILL_HEAD, CENTER, BORDER
ws_i.merge_cells(start_row=r, start_column=3, end_row=r, end_column=4)
r += 1
cond_first = r
for cname, code, how in CONDITIONS:
    cell(ws_i, r, 1, cname, CENTER, fnt=F_BOLD)
    cell(ws_i, r, 2, code, CENTER)
    cell(ws_i, r, 3, how)
    ws_i.merge_cells(start_row=r, start_column=3, end_row=r, end_column=4)
    ws_i.cell(row=r, column=4).border = BORDER
    r += 1
cell(ws_i, r, 1, "统一设置", CENTER, fnt=F_BOLD)
cell(ws_i, r, 2, None)
cell(ws_i, r, 3, "同一个录音 App，16kHz 单声道 WAV；文件名 = 剧本编号-条件代码.wav，例如 G1-S1-Q.wav")
ws_i.merge_cells(start_row=r, start_column=3, end_row=r, end_column=4)
ws_i.cell(row=r, column=4).border = BORDER
r += 2

steps_title = r
ws_i.cell(row=r, column=1, value="使用步骤").font = F_BOLD
r += 1
steps = [
    "1. 在“剧本总览”的“角色分配”栏写清谁演哪个角色。",
    "2. 照“台词”表录音；“动作与环境提示”一栏不念，只用于表演。",
    "3. 允许即兴改口，但录完必须对照录音修改“台词”列，以实际说了什么为准；有效字数和预计时长会自动重算。",
    "4. 每录完一个文件，在“录音计划”登记录音人、日期和实际时长；参考文本校对完成后，再放入共享数据池。",
]
for s in steps:
    ws_i.cell(row=r, column=1, value=s).font = F_BODY
    ws_i.cell(row=r, column=1).alignment = WRAP
    ws_i.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
    r += 1
r += 1

lab_head = r
for c, lab in enumerate(["话术类别", "含义", "", ""], start=1):
    h = ws_i.cell(row=r, column=c, value=lab)
    h.font, h.fill, h.alignment, h.border = F_HEAD, FILL_HEAD, CENTER, BORDER
ws_i.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
r += 1
lab_first = r
for lab, meaning in LABELS:
    cell(ws_i, r, 1, lab, CENTER, fnt=F_BOLD)
    cell(ws_i, r, 2, meaning)
    ws_i.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
    for c in (3, 4):
        ws_i.cell(row=r, column=c).border = BORDER
    r += 1
r += 1

col_head = r
for c, lab in enumerate(["其他列", "含义", "", ""], start=1):
    h = ws_i.cell(row=r, column=c, value=lab)
    h.font, h.fill, h.alignment, h.border = F_HEAD, FILL_HEAD, CENTER, BORDER
ws_i.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
r += 1
cols = [
    ("数字规范写法", "台词里的数字一律按口语读法用汉字写；这一列给出规范写法（如 2800元、15:40、YN-0000-3721），供第 4 组做规范化测评"),
    ("热词", "这句里出现的虚构名称、商品行话和地名，供第 5 组整理热词表、统计专名正确率"),
    ("有效字数", "只数汉字、字母和数字，不数标点和空格；与计算 CER 时的口径一致"),
    ("标注备注", "需要提醒的标注难例，例如第 6 组 S3 中游客复述导游原话的 5 句"),
]
col_first = r
for lab, meaning in cols:
    cell(ws_i, r, 1, lab, CENTER, fnt=F_BOLD)
    cell(ws_i, r, 2, meaning)
    ws_i.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
    for c in (3, 4):
        ws_i.cell(row=r, column=c).border = BORDER
    r += 1
r += 1

leg_head = r
for c, lab in enumerate(["格式", "含义", "", "示例"], start=1):
    h = ws_i.cell(row=r, column=c, value=lab)
    h.font, h.fill, h.alignment, h.border = F_HEAD, FILL_HEAD, CENTER, BORDER
ws_i.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
r += 1
legend = [
    ("浅黄色底", "需要填写的单元格（角色分配、状态、录音人、日期、实际时长、勾选项）", "已录音", FILL_INPUT, F_BODY),
    ("蓝色数字", "可调整的参数（朗读语速、时长目标等）", "220", None, F_INPUT),
    ("黑色数字", "本表内的公式结果，不要直接改写", "4.3", None, F_BODY),
    ("绿色数字", "引用其他工作表的公式", "47", None, F_LINK),
]
leg_first = r
for lab, meaning, sample, fill, fnt in legend:
    cell(ws_i, r, 1, lab, CENTER, fnt=F_BOLD)
    cell(ws_i, r, 2, meaning)
    ws_i.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
    ws_i.cell(row=r, column=3).border = BORDER
    cell(ws_i, r, 4, sample, CENTER, fnt=fnt, fill=fill)
    r += 1
ws_i.cell(row=r, column=1, value="填写示例：角色分配写“导游：张三；游客甲：李四；游客乙：王五；司机：赵六”；状态和勾选项从下拉列表中选择。").font = F_NOTE
ws_i.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
ws_i.cell(row=r, column=1).alignment = WRAP
example_row = r
r += 2

notes_title = r
ws_i.cell(row=r, column=1, value="注意").font = F_BOLD
r += 1
notes = [
    "1. 全部为虚构的教学情景剧。旅行社、店铺、酒店、人名、号码见“虚构名称”表；如与真实企业重名，请替换后再用。",
    "2. 内容只涉及购物、费用、行程和服务态度等商业纠纷，不含辱骂、威胁、政治、民族、宗教或地域内容；施压话术均为温和措辞。",
    "3. 景点介绍只用常识，课前请教师再核对一遍。",
    "4. 预计时长按朗读语速 220 字/分钟估算，可在“剧本总览”顶部修改；实际时长以录音文件为准。",
    "5. 录音只用学生本人的声音，录音前签署同意书；剧本和录音中不出现任何真实个人信息。",
    "6. 第 6 组 S3 中有 5 句是游客复述导游原话，按说话人意图标为“服务态度”，已在“标注备注”列注明；如按话题标注，可改为“威胁消费”。",
]
notes_first = r
for s in notes:
    ws_i.cell(row=r, column=1, value=s).font = F_BODY
    ws_i.cell(row=r, column=1).alignment = WRAP
    ws_i.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
    r += 1

widths(ws_i, [16, 14, 40, 40])
TOTAL_W = 16 + 14 + 40 + 40


def merged_h(row, width):
    text = ws_i.cell(row=row, column=1).value or ""
    n = sum(max(1, math.ceil(disp_len(p) / max((width - 2) * 0.9, 2))) for p in str(text).split("\n"))
    ws_i.row_dimensions[row].height = n * 15 + 6


def merged_h_col(row, col, width):
    text = ws_i.cell(row=row, column=col).value or ""
    n = sum(max(1, math.ceil(disp_len(p) / max((width - 2) * 0.9, 2))) for p in str(text).split("\n"))
    ws_i.row_dimensions[row].height = n * 15 + 6


for rr in range(3, 3 + len(meta)):
    merged_h_col(rr, 2, 14 + 40 + 40)
for rr in range(cond_first, cond_first + len(CONDITIONS) + 1):
    merged_h_col(rr, 3, 80)
for rr in range(steps_title + 1, steps_title + 1 + len(steps)):
    merged_h(rr, TOTAL_W)
for rr in range(lab_first, lab_first + len(LABELS)):
    merged_h_col(rr, 2, 94)
for rr in range(col_first, col_first + len(cols)):
    merged_h_col(rr, 2, 94)
for rr in range(leg_first, leg_first + len(legend)):
    merged_h_col(rr, 2, 54)
merged_h(example_row, TOTAL_W)
for rr in range(notes_first, notes_first + len(notes)):
    merged_h(rr, TOTAL_W)

# ------------------------------------------------------------------ finish
for idx, ws in enumerate(wb.worksheets):
    ws.sheet_view.tabSelected = idx == 0
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
wb.active = 0
wb.save(OUT)
print("saved", OUT, "lines", n_lines, "rows", L_FIRST, L_LAST, "overview", O_FIRST, O_LAST, "record", R_FIRST, R_LAST)
