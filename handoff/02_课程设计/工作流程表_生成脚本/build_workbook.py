import math
import sys
import unicodedata

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT = sys.argv[1]

FONT_NAME = "微软雅黑"
BODY = 10

TEAL = "0F6E56"
TEAL_LIGHT = "E1F5EE"
GRAY_LIGHT = "F1EFE8"
FILL_IN = "FFF2CC"

thin = Side(style="thin", color="B4B2A9")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)


def font(bold=False, size=BODY, color="000000"):
    return Font(name=FONT_NAME, size=size, bold=bold, color=color)


F_BODY = font()
F_BOLD = font(bold=True)
F_HEAD = font(bold=True, color="FFFFFF")
F_TITLE = font(bold=True, size=14, color=TEAL)
F_NOTE = font(size=BODY, color="5F5E5A")
F_INPUT = font(color="0000FF")
F_LINK = font(color="008000")

FILL_HEAD = PatternFill("solid", fgColor=TEAL)
FILL_SUB = PatternFill("solid", fgColor=TEAL_LIGHT)
FILL_TOTAL = PatternFill("solid", fgColor=GRAY_LIGHT)
FILL_INPUT = PatternFill("solid", fgColor=FILL_IN)

WRAP = Alignment(wrap_text=True, vertical="top")
WRAP_CENTER = Alignment(wrap_text=True, vertical="center", horizontal="center")
WRAP_MID = Alignment(wrap_text=True, vertical="center")


def disp_len(text):
    if text is None:
        return 0
    s = str(text)
    if s.startswith("="):
        return 6
    return sum(2 if unicodedata.east_asian_width(ch) in ("W", "F") else 1 for ch in s)


def title(ws, text, note=None, span=6):
    ws["A1"] = text
    ws["A1"].font = F_TITLE
    ws.row_dimensions[1].height = 26
    if note:
        ws["A2"] = note
        ws["A2"].font = F_NOTE
        ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=span)


def header(ws, row, labels, start_col=1):
    for i, lab in enumerate(labels):
        c = ws.cell(row=row, column=start_col + i, value=lab)
        c.font = F_HEAD
        c.fill = FILL_HEAD
        c.alignment = WRAP_CENTER
        c.border = BORDER
    ws.row_dimensions[row].height = 30


def body_cell(ws, row, col, value, align=WRAP, fnt=None, fill=None):
    c = ws.cell(row=row, column=col, value=value)
    c.font = fnt or F_BODY
    c.alignment = align
    c.border = BORDER
    if fill:
        c.fill = fill
    return c


def set_widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def fit_rows(ws, first, last, line_pt=15, pad=6):
    merged_skip = set()
    for rng in ws.merged_cells.ranges:
        if rng.min_row != rng.max_row or rng.min_col != rng.max_col:
            for r in range(rng.min_row, rng.max_row + 1):
                for c in range(rng.min_col, rng.max_col + 1):
                    merged_skip.add((r, c))
    for r in range(first, last + 1):
        lines = 1
        for cell in ws[r]:
            if (r, cell.column) in merged_skip or cell.value is None:
                continue
            width = ws.column_dimensions[get_column_letter(cell.column)].width or 8.43
            usable = max((width - 1.5) * 0.9, 2)
            n = 0
            for part in str(cell.value).split("\n"):
                n += max(1, math.ceil(disp_len(part) / usable))
            lines = max(lines, n)
        ws.row_dimensions[r].height = lines * line_pt + pad


def merged_note_height(ws, row, total_width, line_pt=15, pad=6):
    text = ws.cell(row=row, column=1).value or ""
    n = 0
    for part in str(text).split("\n"):
        n += max(1, math.ceil(disp_len(part) / max(total_width - 2, 2)))
    ws.row_dimensions[row].height = n * line_pt + pad


wb = Workbook()

# ---------------------------------------------------------------- 分组分工
ws_g = wb.active
ws_g.title = "分组分工"

# ---------------------------------------------------------------- 数据需求 (created early so references resolve by name)
ws_d = wb.create_sheet("数据需求")

# Parameter rows on 数据需求 (row numbers referenced by other sheets)
D_STUDENTS = 5
D_GROUPS = 6

# ================================================================ 分组分工 content
title(
    ws_g,
    "分组分工与测评指标",
    "每个人都要先完成“写剧本 → 录音 → 校对参考文本 → 算自己录音的 CER”，专项分工在此之上叠加。"
    "人数列中哪两组为 5 人是建议值，可按实际调整；组长、组员、状态为待填写项。",
    span=10,
)
g_head = ["组", "人数", "虚构剧本场景", "专项", "对应流程步骤", "主要测评指标", "交付物", "组长", "组员", "状态"]
header(ws_g, 4, g_head)
groups = [
    (1, 4, "大巴车上推销土特产", "降噪与语音端点检测", "步骤 2", "加噪前后的 CER；切段是否漏掉人声"),
    (2, 4, "购物店“包厢式”推销", "远距离、手机放口袋时的识别", "步骤 2", "不同录音位置之间的 CER 差距"),
    (3, 5, "擅自变更行程", "说话人分离（区分导游与游客）", "步骤 4", "说话人标错的时长比例"),
    (4, 4, "自费项目费用纠纷", "数字、证号、金额的识别与规范化", "步骤 5", "数字、证号、金额的识别正确率"),
    (5, 4, "兜售物品、索要小费", "热词（虚构的旅行社名、店名、行话）", "步骤 3", "开关热词前后的专名正确率"),
    (6, 4, "言语施压、强迫消费（措辞温和）", "TensorFlow 话术分类（模型 A）", "步骤 6", "各类别的准确率和召回率"),
    (7, 4, "服务态度纠纷", "TensorFlow 话术分类（模型 B）", "步骤 6", "各类别的准确率和召回率，并与第 6 组对比"),
    (8, 5, "正常讲解（对照组，无纠纷）", "时间戳对齐、证据片段导出、误报率测量", "步骤 7", "片段起止时间误差；正常讲解被误标为纠纷的比例"),
]
G_FIRST, G_LAST = 5, 5 + len(groups) - 1
for i, (gid, n, scene, special, step, metric) in enumerate(groups):
    r = G_FIRST + i
    body_cell(ws_g, r, 1, f"第 {gid} 组", WRAP_CENTER)
    body_cell(ws_g, r, 2, n, WRAP_CENTER, fnt=F_INPUT)
    body_cell(ws_g, r, 3, scene)
    body_cell(ws_g, r, 4, special)
    body_cell(ws_g, r, 5, step, WRAP_CENTER)
    body_cell(ws_g, r, 6, metric)
    body_cell(ws_g, r, 7, "基础结果 + 改进后结果 + 对比数据")
    body_cell(ws_g, r, 8, None, fill=FILL_INPUT)
    body_cell(ws_g, r, 9, None, fill=FILL_INPUT)
    body_cell(ws_g, r, 10, "未开始", WRAP_CENTER, fill=FILL_INPUT)
G_TOTAL = G_LAST + 1
body_cell(ws_g, G_TOTAL, 1, "合计", WRAP_CENTER, fnt=F_BOLD, fill=FILL_TOTAL)
body_cell(ws_g, G_TOTAL, 2, f"=SUM(B{G_FIRST}:B{G_LAST})", WRAP_CENTER, fnt=F_BOLD, fill=FILL_TOTAL)
body_cell(
    ws_g,
    G_TOTAL,
    3,
    f"=IF(B{G_TOTAL}='数据需求'!B{D_STUDENTS},\"与班级人数一致\",\"与班级人数不一致，请核对\")",
    WRAP_MID,
    fnt=F_LINK,
    fill=FILL_TOTAL,
)
for c in range(4, 11):
    body_cell(ws_g, G_TOTAL, c, None, fill=FILL_TOTAL)
set_widths(ws_g, [9, 7, 28, 30, 12, 32, 20, 10, 22, 10])
fit_rows(ws_g, G_FIRST, G_TOTAL)
merged_note_height(ws_g, 2, sum([9, 7, 28, 30, 12, 32, 20, 10, 22, 10]))
ws_g.freeze_panes = "B5"

dv_status = DataValidation(type="list", formula1='"未开始,进行中,已完成"', allow_blank=True)
ws_g.add_data_validation(dv_status)
dv_status.add(f"J{G_FIRST}:J{G_LAST}")

# ================================================================ 阶段进度
ws_s = wb.create_sheet("阶段进度", 0)
title(
    ws_s,
    "阶段进度（10 周为基准，附 8 周与 12 周方案）",
    "8 周方案：原理课压缩为 1 周，专项改进缩到 2 周。12 周方案：多出的时间给专项改进和串联测试。"
    "周次为 0 表示开课前。剧本写作和参考文本校对主要在课外完成。",
    span=16,
)
# two-level header
top = ["序号", "阶段", "主要内容", "参与者"]
for i, lab in enumerate(top, start=1):
    c = ws_s.cell(row=4, column=i, value=lab)
    ws_s.merge_cells(start_row=4, start_column=i, end_row=5, end_column=i)
plans = [("10 周方案（基准）", 5), ("8 周方案", 8), ("12 周方案", 11)]
for lab, col in plans:
    ws_s.cell(row=4, column=col, value=lab)
    ws_s.merge_cells(start_row=4, start_column=col, end_row=4, end_column=col + 2)
    for j, sub in enumerate(["起始周", "结束周", "周数"]):
        ws_s.cell(row=5, column=col + j, value=sub)
for i, lab in [(14, "主要产出"), (15, "状态"), (16, "备注")]:
    ws_s.cell(row=4, column=i, value=lab)
    ws_s.merge_cells(start_row=4, start_column=i, end_row=5, end_column=i)
for r in (4, 5):
    for c in range(1, 17):
        cell = ws_s.cell(row=r, column=c)
        cell.font = F_HEAD
        cell.fill = FILL_HEAD
        cell.alignment = WRAP_CENTER
        cell.border = BORDER
ws_s.row_dimensions[4].height = 22
ws_s.row_dimensions[5].height = 22

stages = [
    (
        "课前准备",
        "搭好基础流程（上传、格式统一、识别、时间戳、导出）；准备离线模型、数据池目录和剧本模板；确定统一中间格式，各组模块都按它读入和输出。",
        "教师",
        (0, 0), (0, 0), (0, 0),
        "可运行的基础流程；中间格式说明",
        "可能需要两三周准备时间",
    ),
    (
        "环境与识别基线",
        "搭环境，跑通基础流程，学会算 CER。",
        "全员",
        (1, 2), (1, 2), (1, 2),
        "每人一份 CER 结果",
        "",
    ),
    (
        "原理与 TensorFlow 小实验",
        "讲语谱图、MFCC 等原理；每人动手做一个 TensorFlow 小实验。",
        "全员",
        (3, 4), (3, 3), (3, 4),
        "小实验报告",
        "8 周方案压缩为 1 周",
    ),
    (
        "数据生产",
        "写剧本、分条件录音（安静、嘈杂教室、手机放口袋或远距离）、校对参考文本、标注，汇入共享数据池。",
        "全员",
        (5, 5), (4, 4), (5, 5),
        "共享数据池（教师统一管理）",
        "剧本写作约一周、校对约一周，均在课外；课内录音一次课（4 学时）",
    ),
    (
        "分组专项改进",
        "每组先测一次基础流程的结果，再做改进，然后复测；都用全池数据测，不只测本组录音。",
        "各组",
        (6, 8), (5, 6), (6, 9),
        "各组改进前后的对比数据",
        "8 周方案缩到 2 周；12 周方案延长到 4 周",
    ),
    (
        "串联测试",
        "教师把各组模块接进基础流程，跑全量测试，生成几份核查初稿样例。",
        "教师主导，各组配合",
        (9, 9), (7, 7), (10, 11),
        "核查初稿样例",
        "12 周方案延长到 2 周",
    ),
    (
        "答辩演示",
        "各组演示与答辩。",
        "全员",
        (10, 10), (8, 8), (12, 12),
        "演示与答辩",
        "",
    ),
]
S_FIRST = 6
S_LAST = S_FIRST + len(stages) - 1
for i, (name, content, who, p10, p8, p12, out, note) in enumerate(stages):
    r = S_FIRST + i
    body_cell(ws_s, r, 1, i + 1, WRAP_CENTER)
    body_cell(ws_s, r, 2, name, fnt=F_BOLD)
    body_cell(ws_s, r, 3, content)
    body_cell(ws_s, r, 4, who, WRAP_CENTER)
    for (start, end), col in zip((p10, p8, p12), (5, 8, 11)):
        sc = get_column_letter(col)
        ec = get_column_letter(col + 1)
        body_cell(ws_s, r, col, start, WRAP_CENTER, fnt=F_INPUT)
        body_cell(ws_s, r, col + 1, end, WRAP_CENTER, fnt=F_INPUT)
        body_cell(ws_s, r, col + 2, f'=IF({sc}{r}=0,"开课前",{ec}{r}-{sc}{r}+1)', WRAP_CENTER)
    body_cell(ws_s, r, 14, out)
    body_cell(ws_s, r, 15, "未开始", WRAP_CENTER, fill=FILL_INPUT)
    body_cell(ws_s, r, 16, note or None)
S_TOTAL = S_LAST + 1
body_cell(ws_s, S_TOTAL, 1, None, fill=FILL_TOTAL)
body_cell(ws_s, S_TOTAL, 2, "方案总周数", fnt=F_BOLD, fill=FILL_TOTAL)
for c in (3, 4):
    body_cell(ws_s, S_TOTAL, c, None, fill=FILL_TOTAL)
for col in (5, 8, 11):
    ec = get_column_letter(col + 1)
    body_cell(ws_s, S_TOTAL, col, None, fill=FILL_TOTAL)
    body_cell(ws_s, S_TOTAL, col + 1, None, fill=FILL_TOTAL)
    body_cell(ws_s, S_TOTAL, col + 2, f"=MAX({ec}{S_FIRST}:{ec}{S_LAST})", WRAP_CENTER, fnt=F_BOLD, fill=FILL_TOTAL)
for c in (14, 15, 16):
    body_cell(ws_s, S_TOTAL, c, None, fill=FILL_TOTAL)
set_widths(ws_s, [6, 19, 40, 12, 7, 7, 8, 7, 7, 8, 7, 7, 8, 20, 9, 24])
fit_rows(ws_s, S_FIRST, S_TOTAL)
merged_note_height(ws_s, 2, 120)
ws_s.freeze_panes = "C6"
dv_stage = DataValidation(type="list", formula1='"未开始,进行中,已完成"', allow_blank=True)
ws_s.add_data_validation(dv_stage)
dv_stage.add(f"O{S_FIRST}:O{S_LAST}")

# ================================================================ 处理流程
ws_p = wb.create_sheet("处理流程", 1)
title(
    ws_p,
    "工具处理流程与负责方",
    "每段录音依次经过以下八步。各组只改进自己负责的那一步，输入输出都遵循“统一中间格式”工作表中的字段。",
    span=6,
)
header(ws_p, 4, ["步骤", "名称", "做什么", "负责方", "写入的中间格式字段", "测评指标"])
pipeline = [
    ("上传录音或视频", "统一转为 16kHz 单声道音频", "教师模板", "—", "—"),
    ("降噪与语音端点检测", "降低噪声，切出有人声的段落", "第 1 组；第 2 组负责远距离条件", "开始时间、结束时间", "加噪前后的 CER；是否漏切人声；不同录音位置的 CER 差距"),
    ("语音识别与热词", "把每段转写为文字，加入虚构专名热词", "基线：教师模板；热词：第 5 组", "文字内容", "CER；开关热词前后的专名正确率"),
    ("说话人分离", "区分导游和游客", "第 3 组", "说话人", "说话人标错的时长比例"),
    ("数字、证号、金额规范化", "把口语中的数字、证号、金额转为规范格式", "第 4 组", "文字内容（规范化后）", "数字、证号、金额的识别正确率"),
    ("话术分类（TensorFlow）", "给每段文字打投诉类别标签", "第 6 组、第 7 组（对比两种模型）", "标签", "各类别的准确率和召回率"),
    ("时间戳与证据片段导出", "按时间截取疑似片段并导出", "第 8 组（兼测误报率）", "（输出片段文件）", "片段起止时间误差；正常讲解被误标为纠纷的比例"),
    ("生成核查初稿，人工复核", "汇总为标注“疑似、待核查”的初稿，由人复核", "投诉处理人员（必须）", "—", "—"),
]
P_FIRST = 5
for i, (name, what, owner, field, metric) in enumerate(pipeline):
    r = P_FIRST + i
    is_group = owner.startswith("第") or "：第" in owner
    fill = None if is_group else FILL_TOTAL
    body_cell(ws_p, r, 1, i + 1, WRAP_CENTER, fill=fill)
    body_cell(ws_p, r, 2, name, fnt=F_BOLD, fill=fill)
    body_cell(ws_p, r, 3, what, fill=fill)
    body_cell(ws_p, r, 4, owner, fill=fill)
    body_cell(ws_p, r, 5, field, fill=fill)
    body_cell(ws_p, r, 6, metric, fill=fill)
P_LAST = P_FIRST + len(pipeline) - 1
set_widths(ws_p, [6, 24, 32, 32, 22, 36])
fit_rows(ws_p, P_FIRST, P_LAST)
merged_note_height(ws_p, 2, 152)
ws_p.freeze_panes = "C5"
lg = P_LAST + 2
ws_p.cell(row=lg, column=1, value="灰色底：教师模板或人工环节；白色底：学生小组负责。").font = F_NOTE

# ================================================================ 数据需求 content
ws_d.title = "数据需求"
title(
    ws_d,
    "数据需求估算",
    "录音主要用于测评，不用于从零训练识别模型。蓝色数字为可调整的参数（来自讨论中的建议值），黑色为公式计算结果，绿色为引用其他工作表。",
    span=4,
)
header(ws_d, 4, ["参数", "数值", "单位", "说明"])
params = [
    (D_STUDENTS, "班级人数", 34, "人", "用户提供", F_INPUT),
    (D_GROUPS, "小组数", "=COUNTA('分组分工'!A5:A12)", "组", "自动统计“分组分工”工作表", F_LINK),
    (7, "每组剧本数", 3, "个", "建议值", F_INPUT),
    (8, "每个剧本时长（下限）", 3, "分钟", "建议 3—5 分钟", F_INPUT),
    (9, "每个剧本时长（上限）", 5, "分钟", "建议 3—5 分钟", F_INPUT),
    (10, "录音条件数", 3, "种", "安静、嘈杂教室、手机放口袋或远距离", F_INPUT),
    (11, "话术类别数", 6, "类", "购物安排、费用、行程变更、服务态度、威胁消费、正常讲解", F_INPUT),
    (12, "每类句数（下限）", 150, "句", "建议每类 150—200 句", F_INPUT),
    (13, "每类句数（上限）", 200, "句", "建议每类 150—200 句", F_INPUT),
    (14, "每人编写句数", 30, "句", "每人写 30 句不同说法，可用 AI 改写扩充，须人工审核", F_INPUT),
    (15, "每组精细标注时长", 10, "分钟", "标出谁在什么时间说话，用于测说话人分离", F_INPUT),
]
for r, name, val, unit, note, fnt in params:
    body_cell(ws_d, r, 1, name)
    body_cell(ws_d, r, 2, val, WRAP_CENTER, fnt=fnt)
    body_cell(ws_d, r, 3, unit, WRAP_CENTER)
    body_cell(ws_d, r, 4, note)

header(ws_d, 17, ["计算结果", "下限", "上限", "计算方式"])
results = [
    (18, "每组测评录音（分钟）", "=B7*B8*B10", "=B7*B9*B10", "每组剧本数 × 每个剧本时长 × 录音条件数", "0"),
    (19, "全班测评录音（小时）", "=B18*B6/60", "=C18*B6/60", "每组录音 × 小组数 ÷ 60", "0.0"),
    (20, "话术分类所需句数", "=B11*B12", "=B11*B13", "话术类别数 × 每类句数", "0"),
    (21, "全班可编写句数", "=B14*B5", "=B14*B5", "每人编写句数 × 班级人数", "0"),
    (
        22,
        "句数是否够用",
        '=IF(B21>=B20,"达到下限","未达到下限")',
        '=IF(C21>=C20,"达到上限","未达到上限，可用 AI 改写扩充（须人工审核）")',
        "比较可编写句数与所需句数",
        None,
    ),
    (23, "精细标注总时长（分钟）", "=B15*B6", "=B15*B6", "每组精细标注时长 × 小组数", "0"),
]
for r, name, lo, hi, how, fmt in results:
    body_cell(ws_d, r, 1, name)
    c1 = body_cell(ws_d, r, 2, lo, WRAP_CENTER)
    c2 = body_cell(ws_d, r, 3, hi, WRAP_CENTER)
    if fmt:
        c1.number_format = fmt
        c2.number_format = fmt
    body_cell(ws_d, r, 4, how)

header(ws_d, 25, ["环节", "形式", "时间", "说明"])
timing = [
    (26, "剧本写作", "课外", "约一周", "名字一律虚构；只写商业纠纷，不写侮辱性、地域或民族类言论"),
    (27, "录音", "课内", "一次课（4 学时）", "统一录音 App，16kHz 单声道，统一命名规则；允许即兴、打断、带情绪"),
    (28, "校对参考文本", "课外", "约一周", "对照录音修正，以实际说了什么为准"),
]
for r, a, b, c, d in timing:
    body_cell(ws_d, r, 1, a)
    body_cell(ws_d, r, 2, b, WRAP_CENTER)
    body_cell(ws_d, r, 3, c, WRAP_CENTER)
    body_cell(ws_d, r, 4, d)
set_widths(ws_d, [24, 14, 22, 52])
fit_rows(ws_d, 5, 15)
fit_rows(ws_d, 18, 23)
fit_rows(ws_d, 26, 28)
merged_note_height(ws_d, 2, 112)

# ================================================================ 统一中间格式
ws_m = wb.create_sheet("统一中间格式")
title(
    ws_m,
    "统一中间格式",
    "每段文字都记录以下五个字段，各组模块按这个格式读入和输出，这样可以独立开发、最后串联。示例列仅用于说明格式。",
    span=4,
)
header(ws_m, 4, ["字段", "含义", "示例", "由哪一步写入"])
fields = [
    ("开始时间（秒）", "该段在录音中的起点", "12.4", "步骤 2、3"),
    ("结束时间（秒）", "该段在录音中的终点", "18.9", "步骤 2、3"),
    ("说话人", "导游、游客或未知", "导游", "步骤 4"),
    ("文字内容", "识别并规范化后的文字", "这个手镯今天优惠价 2800 元", "步骤 3、5"),
    ("标签", "投诉类别，一律标为“疑似”", "疑似·费用", "步骤 6"),
]
for i, row in enumerate(fields):
    r = 5 + i
    for j, v in enumerate(row, start=1):
        body_cell(ws_m, r, j, v, WRAP_CENTER if j in (3, 4) and j != 3 else WRAP, fnt=F_BOLD if j == 1 else None)
set_widths(ws_m, [18, 30, 30, 16])
fit_rows(ws_m, 5, 9)
merged_note_height(ws_m, 2, 90)

# ================================================================ 协作规则
ws_r = wb.create_sheet("协作规则")
title(ws_r, "贯穿全程的规则", None)
header(ws_r, 3, ["序号", "规则"])
rules = [
    "每个人都要完成“写剧本、录音、校对参考文本、算自己录音的 CER”，专项分工是在这之上叠加的。",
    "每组的交付都是“基础结果 + 改进后结果 + 对比数据”，不要求在真实场景中好用。",
    "核查初稿一律标为“疑似、待核查”，最后一步人工复核不能省。",
    "各组模块都按“统一中间格式”读入和输出，可以独立开发，最后由教师串联。",
    "录音和标注全部进入共享数据池，由教师统一管理；各组专项都用全池数据测试，不只测本组录音。",
]
for i, text in enumerate(rules):
    r = 4 + i
    body_cell(ws_r, r, 1, i + 1, WRAP_CENTER)
    body_cell(ws_r, r, 2, text)
set_widths(ws_r, [6, 90])
fit_rows(ws_r, 4, 3 + len(rules))

# ================================================================ 说明
ws_i = wb.create_sheet("说明", 0)
ws_i["A1"] = "语音识别实训课 · 旅游纠纷录音材料整理项目：工作流程与分组"
ws_i["A1"].font = F_TITLE
ws_i.row_dimensions[1].height = 26
meta = [
    ("整理日期", "2026-10-07"),
    ("内容来源", "本次对话中的课程设计讨论（方案草稿，尚未定稿）"),
    ("适用对象", "约 34 名专科生，4—5 人一组，共 8 组；8—12 周"),
]
for i, (k, v) in enumerate(meta):
    r = 3 + i
    body_cell(ws_i, r, 1, k, fnt=F_BOLD, fill=FILL_SUB)
    body_cell(ws_i, r, 2, v)
    ws_i.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
    for c in (3,):
        ws_i.cell(row=r, column=c).border = BORDER

header(ws_i, 7, ["工作表", "内容", ""])
ws_i.merge_cells(start_row=7, start_column=2, end_row=7, end_column=3)
sheets = [
    ("阶段进度", "七个阶段的内容、参与者和产出；同时列出 10 周、8 周、12 周三种排法"),
    ("处理流程", "工具的八个处理步骤、负责方、写入字段和测评指标"),
    ("分组分工", "8 个组的虚构剧本场景、专项、测评指标、交付物，以及待填写的组长、组员和状态"),
    ("数据需求", "测评录音、话术文本、精细标注的数量估算，可调参数后自动重算"),
    ("统一中间格式", "各组模块之间传递数据的五个字段"),
    ("协作规则", "贯穿全程的五条规则"),
]
for i, (name, desc) in enumerate(sheets):
    r = 8 + i
    body_cell(ws_i, r, 1, name, fnt=F_BOLD)
    body_cell(ws_i, r, 2, desc)
    ws_i.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
    ws_i.cell(row=r, column=3).border = BORDER

lg = 8 + len(sheets) + 1
header(ws_i, lg, ["格式", "含义", "示例"])
legend = [
    ("浅黄色底", "需要填写的单元格（组长、组员、状态）", "进行中", FILL_INPUT, F_BODY),
    ("蓝色数字", "可调整的参数，改动后相关公式会自动重算", "34", None, F_INPUT),
    ("黑色数字", "公式计算结果，不要直接改写", "3.6", None, F_BODY),
    ("绿色", "引用其他工作表的公式", "8", None, F_LINK),
    ("灰色底", "合计行，或教师模板、人工环节", "", FILL_TOTAL, F_BODY),
]
for i, (fmt_name, meaning, sample, fill, fnt) in enumerate(legend):
    r = lg + 1 + i
    body_cell(ws_i, r, 1, fmt_name, fnt=F_BOLD)
    body_cell(ws_i, r, 2, meaning)
    body_cell(ws_i, r, 3, sample, WRAP_CENTER, fnt=fnt, fill=fill)

ex = lg + 1 + len(legend) + 1
ws_i.cell(row=ex, column=1, value="填写示例").font = F_BOLD
ws_i.cell(row=ex + 1, column=1, value="组长填“张三”；组员填“李四、王五、赵六”（用顿号分隔）；状态从下拉列表中选择。").font = F_BODY
ws_i.merge_cells(start_row=ex + 1, start_column=1, end_row=ex + 1, end_column=3)
ws_i.cell(row=ex + 1, column=1).alignment = WRAP

nt = ex + 3
ws_i.cell(row=nt, column=1, value="注意").font = F_BOLD
notes = (
    "1. 数据量和周数都是讨论中的估算与建议值，开课前需结合机房条件实测调整。\n"
    "2. 剧本中的旅行社、购物店、导游名字一律虚构；只写商业纠纷，不写政治、民族或侮辱性言论。\n"
    "3. 工具输出只作“疑似、待核查”的提示，不作定性结论，必须人工复核。"
)
ws_i.cell(row=nt + 1, column=1, value=notes).font = F_BODY
ws_i.merge_cells(start_row=nt + 1, start_column=1, end_row=nt + 1, end_column=3)
ws_i.cell(row=nt + 1, column=1).alignment = WRAP

set_widths(ws_i, [16, 56, 16])
fit_rows(ws_i, 3, 5)
for r in range(8, 8 + len(sheets)):
    n = math.ceil(disp_len(ws_i.cell(row=r, column=2).value) / (56 + 16 - 2))
    ws_i.row_dimensions[r].height = max(1, n) * 15 + 6
fit_rows(ws_i, lg + 1, lg + len(legend))
merged_note_height(ws_i, ex + 1, 88)
merged_note_height(ws_i, nt + 1, 88)

# ---------------------------------------------------------------- final sheet order and print setup
order = ["说明", "阶段进度", "处理流程", "分组分工", "数据需求", "统一中间格式", "协作规则"]
wb._sheets = [wb[name] for name in order]
for idx, ws in enumerate(wb.worksheets):
    ws.sheet_view.tabSelected = idx == 0
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
wb.active = 0

wb.save(OUT)
print("saved", OUT)
