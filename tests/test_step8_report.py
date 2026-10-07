"""步骤 8 核查初稿（pipeline/step8_report.py）的测试。

- Word 初稿：标题、醒目声明、文件信息、摘要（各类疑似片段数、复核状态）、疑似片段表、
  全文时间轴转写、签字栏；文档属性写明由人工智能技术自动生成；不出现禁用说法；
- 导出 ZIP：review_draft.docx、segments.csv、segments.json、clips/、说明.txt 五类文件，
  说明.txt 和 JSON 的 meta.notice 都是通知语；ZIP 文件名只用 ASCII。
"""
import json
import zipfile
from pathlib import Path

import numpy as np
import pytest

from pipeline.audio import SR
from pipeline.data import FLAG_OUTPUTS
from pipeline.schema import NOTICE, new_segment, read_csv
from pipeline.step7_clips import export_clips
from pipeline.step8_report import (
    AUTHOR,
    DISCLAIMER,
    FORBIDDEN_WORDS,
    TITLE,
    build_docx,
    export_bundle,
    fmt_mmss,
)

docx = pytest.importorskip("docx")


def _segments() -> list[dict]:
    """示例段落：3 段有标签（其中 1 段是消费施压），2 段没有；第 1 段已人工确认。"""
    return [
        new_segment(12.4, 18.9, speaker="导游", text="这个手镯今天优惠价2800元。", label="疑似·费用",
                    category="费用", numbers=["2800元"], entities=["雾隐行舟旅行社"],
                    review="确认", review_note="已听原声"),
        new_segment(20.0, 25.55, speaker="游客", text="电话是0871-0000-6688。", numbers=["0871-0000-6688"]),
        new_segment(30.0, 36.0, speaker="导游", text="不进店的话，晚上的住宿就不好安排了。", label="疑似·消费施压",
                    category="威胁消费"),
        new_segment(40.0, 44.0, speaker="导游", text="下午3点我们去石林。", label="疑似·行程变更",
                    category="行程变更", numbers=["15:00"]),
        new_segment(125.4, 130.0, speaker="导游", text="大家注意脚下安全。", category="正常讲解"),
    ]


def _meta() -> dict:
    return {
        "file": "G1-S1-Q.m4a",
        "duration": 268.4,
        "sha256": "a" * 64,
        "processed_at": "2026-10-07T15:20:00+08:00",
        "models": {"asr": "sense-voice-int8-2024-07-17", "vad": "silero_vad",
                   "diarization": "pyannote-3.0 + campplus"},
        "options": {"num_speakers": 2, "hotword_fix": False},
        "methods": {"denoise": "baseline", "classify": "baseline"},
        "notice": NOTICE,
    }


def _all_text(document) -> str:
    """Word 文档里能看到的全部文字：正文段落、表格、页脚、文档属性。"""
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    for section in document.sections:
        parts.extend(p.text for p in section.footer.paragraphs)
        parts.extend(p.text for p in section.header.paragraphs)
    props = document.core_properties
    parts.extend([props.title or "", props.author or "", props.comments or "", props.subject or "",
                  props.keywords or ""])
    return "\n".join(parts)


def test_constants_verbatim():
    assert TITLE == "旅游纠纷录音材料核查初稿（疑似、待核查）"
    assert DISCLAIMER == ("本初稿由语音识别等人工智能技术自动生成，识别和分类都可能出错。所有标注均为“疑似、待核查”，"
                          "不代表任何定性结论，必须由工作人员对照原始录音逐条复核。本工具不鉴定录音的真伪。")
    assert AUTHOR == "旅游纠纷录音材料整理工具（教学原型）"
    assert len(FORBIDDEN_WORDS) == 4


@pytest.mark.parametrize("seconds, text", [
    (125.4, "02:05.4"),
    (0, "00:00.0"),
    (9.94, "00:09.9"),
    (59.96, "01:00.0"),
    (3725.0, "62:05.0"),
])
def test_fmt_mmss(seconds, text):
    assert fmt_mmss(seconds) == text


def test_docx_contents(tmp_path):
    path = tmp_path / "review_draft.docx"
    build_docx(_segments(), _meta(), path)
    document = docx.Document(str(path))

    paragraphs = [p.text for p in document.paragraphs]
    assert any(TITLE in text for text in paragraphs)
    assert any(DISCLAIMER in text for text in paragraphs)
    # 标题在声明前面，声明在文件信息前面（build_spec 5.8 的顺序）
    title_at = next(i for i, t in enumerate(paragraphs) if TITLE in t)
    disclaimer_at = next(i for i, t in enumerate(paragraphs) if DISCLAIMER in t)
    assert title_at < disclaimer_at

    props = document.core_properties
    assert props.author == AUTHOR
    assert props.title == TITLE
    assert props.comments == "由人工智能技术自动生成"

    full = _all_text(document)
    for word in FORBIDDEN_WORDS:
        assert word not in full
    # 文件信息
    assert "G1-S1-Q.m4a" in full and "a" * 64 in full and "sense-voice-int8-2024-07-17" in full
    # 全文时间轴转写里有每一段（时间用 分:秒）
    assert "02:05.4" in full and "大家注意脚下安全。" in full
    # 签字栏
    assert "复核人" in full and "复核日期" in full


def test_docx_section_order(tmp_path):
    path = tmp_path / "review_draft.docx"
    build_docx(_segments(), _meta(), path)
    paragraphs = [p.text for p in docx.Document(str(path)).paragraphs]
    headings = ["一、文件信息", "二、摘要", "三、疑似片段", "四、全文时间轴转写", "五、人工复核签字"]
    positions = [paragraphs.index(h) for h in headings]
    assert positions == sorted(positions)


def test_docx_summary_counts_per_label(tmp_path):
    path = tmp_path / "review_draft.docx"
    build_docx(_segments(), _meta(), path)
    full = _all_text(docx.Document(str(path)))
    # 每个疑似标签都列出段数（没有的写 0 段）
    assert "疑似·费用：1 段" in full
    assert "疑似·消费施压：1 段" in full
    assert "疑似·行程变更：1 段" in full
    assert "疑似·购物安排：0 段" in full
    assert "疑似·服务态度：0 段" in full
    # 复核状态：疑似片段 3 段里 1 段已复核
    assert "已复核 1 条 / 共 3 条" in full


def test_docx_suspicious_table(tmp_path):
    path = tmp_path / "review_draft.docx"
    build_docx(_segments(), _meta(), path)
    document = docx.Document(str(path))
    table = next(t for t in document.tables if t.rows[0].cells[0].text == "序号")
    header = [c.text for c in table.rows[0].cells]
    assert header == ["序号", "起止时间", "说话人", "文字", "标签", "数字和名称", "复核结论", "复核意见"]
    rows = [[c.text for c in row.cells] for row in table.rows[1:]]
    assert len(rows) == 3  # 只有带标签的段
    assert rows[0] == ["1", "00:12.4–00:18.9", "导游", "这个手镯今天优惠价2800元。", "疑似·费用",
                       "2800元；雾隐行舟旅行社", "确认", "已听原声"]
    assert rows[1][0] == "3" and rows[1][4] == "疑似·消费施压" and rows[1][6] == "未复核"


def test_docx_chinese_font(tmp_path):
    """正文样式设了中文字体（eastAsia），Windows 上的 Word 才能正常显示中文。"""
    from docx.oxml.ns import qn

    path = tmp_path / "review_draft.docx"
    build_docx(_segments(), _meta(), path)
    document = docx.Document(str(path))
    for name in ("Normal", "Title", "Heading 1"):
        fonts = document.styles[name].element.rPr.rFonts
        assert fonts.get(qn("w:eastAsia")), name
        assert fonts.get(qn("w:eastAsiaTheme")) is None, name  # 主题字体会盖过 eastAsia，必须去掉


def test_docx_empty(tmp_path):
    """没有段落、meta 也是空的（例如录音里没有人声）：照样生成，声明还在。"""
    path = tmp_path / "empty.docx"
    build_docx([], {}, path)
    full = _all_text(docx.Document(str(path)))
    assert DISCLAIMER in full
    assert "没有被标出的疑似片段" in full


def _tone(seconds: float) -> np.ndarray:
    t = np.arange(int(seconds * SR)) / SR
    return (0.1 * np.sin(2 * np.pi * 440 * t)).astype(np.float32)


def test_bundle_zip(tmp_path):
    segments = _segments()
    clips_dir = tmp_path / "work" / "clips"
    segments = export_clips(segments, _tone(140.0), SR, clips_dir, {"clips": {"padding": 1.0}})
    out_dir = tmp_path / "out"
    zip_path = export_bundle(segments, _meta(), out_dir, clips_dir=str(clips_dir))

    assert isinstance(zip_path, str)
    assert Path(zip_path) == out_dir / "G1-S1-Q_review.zip"
    with zipfile.ZipFile(zip_path) as zf:
        names = zf.namelist()
        assert "review_draft.docx" in names
        assert "segments.csv" in names
        assert "segments.json" in names
        assert "说明.txt" in names
        clip_names = [n for n in names if n.startswith("clips/") and n.endswith(".wav")]
        assert len(clip_names) == 3
        assert "clips/clips_index.csv" in names

        readme = zf.read("说明.txt").decode("utf-8-sig")
        assert NOTICE in readme
        assert "人工智能技术自动生成" in readme
        for word in FORBIDDEN_WORDS:
            assert word not in readme

        data = json.loads(zf.read("segments.json").decode("utf-8"))
        assert data["meta"]["notice"] == NOTICE
        assert "人工智能技术自动生成" in data["meta"]["generated_by"]
        assert data["segments"][0]["review"] == "确认"

    # Word、CSV、JSON 也留在 out_dir 里，界面可以单独下载
    assert (out_dir / "review_draft.docx").exists()
    assert read_csv(out_dir / "segments.csv")[0]["speaker"] == "导游"
    assert (out_dir / "segments.json").exists()


def test_bundle_without_clips_and_non_ascii_name(tmp_path):
    meta = _meta()
    meta["file"] = "张三的录音.m4a"
    zip_path = export_bundle(_segments(), meta, tmp_path, clips_dir=None)
    name = Path(zip_path).name
    assert name.isascii() and name.endswith("_review.zip")
    assert name == "______review.zip"  # 5 个汉字各换成 _，再接 _review.zip
    with zipfile.ZipFile(zip_path) as zf:
        names = zf.namelist()
        assert {"review_draft.docx", "segments.csv", "segments.json", "说明.txt"} <= set(names)
        assert any(n.startswith("clips/") for n in names)  # 没有片段音频时也留一个 clips/ 文件夹
        info = zf.getinfo("说明.txt")
        assert info.flag_bits & 0x800  # 文件名按 UTF-8 存，解压时中文名不乱码


def test_bundle_does_not_change_input(tmp_path):
    segments = _segments()
    meta = {"file": "demo.wav"}
    before = json.dumps([segments, meta], ensure_ascii=False, sort_keys=True)
    export_bundle(segments, meta, tmp_path)
    assert json.dumps([segments, meta], ensure_ascii=False, sort_keys=True) == before


def test_labels_used_are_legal():
    """示例数据里用到的标签都在合法范围内（防止测试本身写错标签）。"""
    assert {seg["label"] for seg in _segments()} <= {""} | set(FLAG_OUTPUTS)
