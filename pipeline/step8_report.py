"""步骤 8：核查初稿（Word）与导出压缩包（教师模板，学生不改）。

把整理好的段落写成一份给工作人员人工复核用的 Word 初稿，再连同表格、片段音频打成一个 ZIP。
工具只提示、不判定：初稿里所有标注都是"疑似、待核查"，必须由人对照原始录音逐条复核。

Word 初稿的内容依次是（照 docs/build_spec.md 5.8）：
    1. 标题 TITLE；
    2. 醒目的声明 DISCLAIMER（红色加粗）；
    3. 一、文件信息：原始文件名、时长、SHA-256、处理时间、模型名称和版本、处理参数、各步骤做法；
    4. 二、摘要：各类疑似片段的数量、复核状态（已复核 N 条 / 共 M 条）；
    5. 三、疑似片段：序号、起止时间（分:秒）、说话人、文字、标签、数字和名称、复核结论、复核意见；
    6. 四、全文时间轴转写：每段一行；
    7. 五、人工复核签字：复核人、复核日期。
    每页页脚都印着通知语 NOTICE。文档属性（文件 → 信息里能看到）写明由人工智能技术自动生成。

导出的 ZIP（export_bundle）里有：
    review_draft.docx  核查初稿
    segments.csv       段落表格（UTF-8 带 BOM，中文表头，Excel 双击打开不乱码）
    segments.json      段落和处理信息（meta.notice 是通知语）
    clips/             疑似片段音频和索引表 clips_index.csv（步骤 7 导出的）
    说明.txt           写明由人工智能技术自动生成、识别可能有误、必须人工复核

基线做法：
    用 python-docx 按上面的顺序写 Word；正文用宋体、标题用微软雅黑（设置中文字体 eastAsia，
    Windows 上的 Word 才不会把中文显示成别的字体）；纸张 A4。
    用标准库 zipfile 打包，ZIP 里的中文文件名按 UTF-8 存。
可改进方向：
    这是教师模板，学生不改。界面上的表格改了说话人、文字、标签、复核结论后再导出，
    导出用的就是改过的段落（见 pipeline/table.py 的 rows_to_segments）。
测评指标：
    不涉及识别效果；由 tests/test_step8_report.py 检查内容顺序、声明、文档属性、
    不出现禁用的说法、ZIP 里五类文件齐全。
"""
import zipfile
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from pipeline.data import FLAG_OUTPUTS
from pipeline.methods import SLOT_TITLES
from pipeline.schema import GENERATED_BY, LIST_SEP, NOTICE, REVIEW_CHOICES, write_csv, write_json
from pipeline.step7_clips import CLIPS_INDEX_NAME

TITLE = "旅游纠纷录音材料核查初稿（疑似、待核查）"
DISCLAIMER = "本初稿由语音识别等人工智能技术自动生成，识别和分类都可能出错。所有标注均为“疑似、待核查”，不代表任何定性结论，必须由工作人员对照原始录音逐条复核。本工具不鉴定录音的真伪。"
AUTHOR = "旅游纠纷录音材料整理工具（教学原型）"

# 界面、初稿、导出文件里都不能出现的说法（红线第 4 条）。这一行以外，pipeline/ 里不要再写这些词
FORBIDDEN_WORDS = ["违规", "违法", "执法级准确率", "可作为法律证据"]

# Word 文档属性"备注"（隐式标识：说明这份文件是机器生成的）
DOC_COMMENTS = "由人工智能技术自动生成"

# 导出文件的名字（ZIP 里也是这些名字）
DOCX_NAME = "review_draft.docx"
CSV_NAME = "segments.csv"
JSON_NAME = "segments.json"
README_NAME = "说明.txt"
CLIPS_FOLDER = "clips"

# 字体：正文宋体，标题微软雅黑（中文 Windows 都自带）
BODY_FONT = "宋体"
HEADING_FONT = "微软雅黑"

# 疑似片段表的表头
SUSPECT_HEADERS = ["序号", "起止时间", "说话人", "文字", "标签", "数字和名称", "复核结论", "复核意见"]
SUSPECT_WIDTHS_CM = [1.0, 2.6, 1.6, 4.6, 2.2, 2.4, 1.4, 2.0]

# meta 里英文键的中文名（文件信息表里用）
MODEL_TITLES = {"asr": "语音识别", "vad": "端点检测", "diarization": "说话人分离"}
OPTION_TITLES = {
    "num_speakers": "说话人数（-1 为自动）",
    "denoise": "降噪",
    "hotword_fix": "热词纠错",
    "hotwords": "热词表",
    "methods": "做法",
    "asr_mode": "识别模式",
    "out_dir": "输出文件夹",
}

# 没有提供的信息显示成这样
MISSING = "（未提供）"


# ---------------- 小工具 ----------------

def fmt_mmss(seconds) -> str:
    """秒 → "分:秒"，秒保留 1 位小数。例：125.4 → "02:05.4"；3725 → "62:05.0"。"""
    tenths = max(0, int(round(float(seconds or 0) * 10)))  # 先换成"多少个 0.1 秒"，避免 59.96 显示成 00:60.0
    minutes, tenths = divmod(tenths, 600)
    return f"{minutes:02d}:{tenths / 10:04.1f}"


def _as_list(value) -> list:
    """numbers、entities 可能是列表，也可能是界面里改过的文字（"2800元；15:40"），统一成列表。"""
    if not value:
        return []
    if isinstance(value, str):
        return [value]
    return list(value)


def _short(value) -> str:
    """把配置里的值写成短短一句：开/关、列表只列前 10 个。"""
    if isinstance(value, bool):
        return "开" if value else "关"
    if isinstance(value, (list, tuple)):
        text = "、".join(str(v) for v in value[:10])
        return text + (f"…（共 {len(value)} 个）" if len(value) > 10 else "")
    if isinstance(value, dict):
        return "，".join(f"{k}={v}" for k, v in value.items())
    if value is None:
        return "（未设置）"
    return str(value)


def _format_mapping(mapping, titles: dict) -> str:
    """{"asr": "sense-voice…", …} → "语音识别：sense-voice…；…"；空的返回 ""。"""
    if not isinstance(mapping, dict) or not mapping:
        return ""
    return LIST_SEP.join(f"{titles.get(key, key)}：{_short(value)}" for key, value in mapping.items())


def _is_reviewed(segment: dict) -> bool:
    """复核结论是确认、修改、驳回之一，就算已复核。"""
    return segment.get("review") in REVIEW_CHOICES[1:]


# ---------------- Word 的格式 ----------------

def _set_style_font(style, font_name: str) -> None:
    """给一个样式设字体：西文字体和中文字体（eastAsia）都用 font_name。

    样式里原来写的"主题字体"（asciiTheme、eastAsiaTheme 等）优先级更高，会盖过我们设的字体，所以要删掉。
    """
    style.font.name = font_name  # 西文字体
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:eastAsia"), font_name)  # 中文字体
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        fonts.attrib.pop(qn(attr), None)


def _setup_document(doc) -> None:
    """纸张 A4、页边距 2 厘米、中文字体、页脚通知语、文档属性。"""
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21.0), Cm(29.7)
    section.left_margin = section.right_margin = Cm(2.0)
    section.top_margin = section.bottom_margin = Cm(2.0)
    section.footer.paragraphs[0].text = f"{NOTICE}（{DOC_COMMENTS}）"

    _set_style_font(doc.styles["Normal"], BODY_FONT)
    doc.styles["Normal"].font.size = Pt(10.5)  # 五号字
    for name in ("Title", "Heading 1", "Heading 2"):
        _set_style_font(doc.styles[name], HEADING_FONT)
    doc.styles["Title"].font.size = Pt(20)  # 标题一行放得下

    props = doc.core_properties
    props.title = TITLE
    props.author = AUTHOR
    props.last_modified_by = AUTHOR
    props.comments = DOC_COMMENTS
    props.subject = NOTICE
    now = datetime.now(timezone.utc)
    props.created = now
    props.modified = now


def _fill_cells(cells, values, bold: bool = False, size: float = 9) -> None:
    """把一行文字写进表格的格子里（小号字，表头加粗）。"""
    for cell, value in zip(cells, values):
        run = cell.paragraphs[0].add_run(str(value))
        run.bold = bold
        run.font.size = Pt(size)


def _set_widths(table, widths_cm) -> None:
    """设置每一列的宽度（Word 要求每个格子都设一遍）。"""
    for row in table.rows:
        for cell, width in zip(row.cells, widths_cm):
            cell.width = Cm(width)


# ---------------- Word 的各部分 ----------------

def _file_info_rows(meta: dict) -> list[tuple[str, str]]:
    """文件信息表的每一行：(项目, 内容)。"""
    duration = meta.get("duration")
    rows = [
        ("原始文件", meta.get("file")),
        ("时长", f"{fmt_mmss(duration)}（{float(duration):.1f} 秒）" if duration is not None else ""),
        ("SHA-256", meta.get("sha256")),
        ("处理时间", meta.get("processed_at")),
        ("模型", _format_mapping(meta.get("models"), MODEL_TITLES)),
        ("处理参数", _format_mapping(meta.get("options"), OPTION_TITLES)),
        ("各步骤做法", _format_mapping(meta.get("methods"), SLOT_TITLES)),
        ("生成方式", f"{DOC_COMMENTS}（{AUTHOR}）"),
    ]
    if meta.get("demo_notice"):
        rows.append(("说明", meta["demo_notice"]))
    if meta.get("message"):
        rows.append(("提示", meta["message"]))
    if meta.get("qc"):
        rows.append(("录音质检", LIST_SEP.join(_as_list(meta["qc"]))))
    return [(key, str(value) if value not in (None, "") else MISSING) for key, value in rows]


def _add_file_info(doc, meta: dict) -> None:
    doc.add_heading("一、文件信息", level=1)
    rows = _file_info_rows(meta)
    table = doc.add_table(rows=len(rows), cols=2)
    table.style = "Table Grid"
    for row, (key, value) in zip(table.rows, rows):
        _fill_cells(row.cells[:1], [key], bold=True)
        _fill_cells(row.cells[1:], [value])
    _set_widths(table, [3.0, 14.0])


def _add_summary(doc, segments: list[dict], flagged: list[tuple[int, dict]]) -> None:
    doc.add_heading("二、摘要", level=1)
    doc.add_paragraph(f"全部段落 {len(segments)} 段，其中疑似片段 {len(flagged)} 段。各类疑似片段数量：")
    counts = {label: 0 for label in FLAG_OUTPUTS}
    for _, segment in flagged:
        counts[segment["label"]] = counts.get(segment["label"], 0) + 1
    for label, count in counts.items():
        doc.add_paragraph(f"{label}：{count} 段", style="List Bullet")
    reviewed_flagged = sum(1 for _, segment in flagged if _is_reviewed(segment))
    reviewed_all = sum(1 for segment in segments if _is_reviewed(segment))
    doc.add_paragraph(
        f"复核状态：疑似片段已复核 {reviewed_flagged} 条 / 共 {len(flagged)} 条；"
        f"全部段落已复核 {reviewed_all} 条 / 共 {len(segments)} 条"
        f"（复核结论为“确认”“修改”“驳回”的算已复核）。"
    )


def _add_suspect_table(doc, flagged: list[tuple[int, dict]]) -> None:
    doc.add_heading("三、疑似片段", level=1)
    if not flagged:
        doc.add_paragraph("没有被标出的疑似片段。")
        return
    doc.add_paragraph("序号与全文时间轴转写、表格文件里的序号一致。")
    table = doc.add_table(rows=1, cols=len(SUSPECT_HEADERS))
    table.style = "Table Grid"
    _fill_cells(table.rows[0].cells, SUSPECT_HEADERS, bold=True)
    for n, segment in flagged:
        values = [
            n,
            f"{fmt_mmss(segment.get('start'))}–{fmt_mmss(segment.get('end'))}",
            segment.get("speaker", "未知"),
            segment.get("text", ""),
            segment.get("label", ""),
            LIST_SEP.join(str(v) for v in _as_list(segment.get("numbers")) + _as_list(segment.get("entities"))),
            segment.get("review") or REVIEW_CHOICES[0],
            segment.get("review_note", ""),
        ]
        _fill_cells(table.add_row().cells, values)
    _set_widths(table, SUSPECT_WIDTHS_CM)


def _add_transcript(doc, segments: list[dict]) -> None:
    doc.add_heading("四、全文时间轴转写", level=1)
    if not segments:
        doc.add_paragraph("没有识别出文字。")
        return
    for n, segment in enumerate(segments, start=1):
        para = doc.add_paragraph()
        time_run = para.add_run(f"[{n}] {fmt_mmss(segment.get('start'))}–{fmt_mmss(segment.get('end'))}　")
        time_run.font.color.rgb = RGBColor(0x59, 0x59, 0x59)  # 灰色
        para.add_run(f"{segment.get('speaker', '未知')}：").bold = True
        para.add_run(segment.get("text", ""))
        if segment.get("label"):
            label_run = para.add_run(f"　【{segment['label']}】")
            label_run.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)  # 红色


def _add_signature(doc) -> None:
    doc.add_heading("五、人工复核签字", level=1)
    doc.add_paragraph("本人已对照原始录音逐条复核以上内容，复核结论见“三、疑似片段”表。")
    doc.add_paragraph("复核人（签字）：＿＿＿＿＿＿＿＿　　　复核日期：＿＿＿＿年＿＿月＿＿日")


def build_docx(segments: list[dict], meta: dict, path) -> None:
    """生成 Word 核查初稿，写到 path。内容顺序见本文件开头的说明。"""
    meta = meta or {}
    doc = Document()
    _setup_document(doc)

    # 1. 标题
    doc.add_heading(TITLE, level=0)
    # 2. 醒目的声明：红色加粗
    run = doc.add_paragraph().add_run(DISCLAIMER)
    run.bold = True
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)

    # 疑似片段：(序号, 段落)，序号从 1 开始，是段落在全部段落里的位置
    flagged = [(n, segment) for n, segment in enumerate(segments, start=1) if segment.get("label")]
    _add_file_info(doc, meta)  # 3
    _add_summary(doc, segments, flagged)  # 4
    _add_suspect_table(doc, flagged)  # 5
    _add_transcript(doc, segments)  # 6
    _add_signature(doc)  # 7

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))


# ---------------- 导出压缩包 ----------------

def _ascii_stem(file_name) -> str:
    """原文件名去掉扩展名，非 ASCII 字符（如中文）和 Windows 文件名里不能用的字符换成 _。"""
    stem = Path(str(file_name or "")).stem
    stem = "".join(ch if ch.isascii() and ch.isprintable() and ch not in '\\/:*?"<>|' else "_" for ch in stem)
    return stem or "result"


def _readme_text(meta: dict, clip_count: int) -> str:
    """说明.txt 的内容（用 Windows 换行，记事本打开不挤成一行）。"""
    if clip_count:
        clips_line = f"  clips/              疑似片段音频（WAV，共 {clip_count} 个）和索引表 {CLIPS_INDEX_NAME}"
    else:
        clips_line = "  clips/              疑似片段音频（本次没有：没有被标出的段落，或者没有录音，如剧本文本演示）"
    lines = [
        "旅游纠纷录音材料核查初稿 · 说明",
        "",
        f"本压缩包{DOC_COMMENTS}（{AUTHOR}）。",
        f"{NOTICE}。",
        DISCLAIMER,
        "",
        f"原始文件：{meta.get('file') or MISSING}",
        f"原始文件 SHA-256：{meta.get('sha256') or MISSING}",
        f"导出时间：{meta.get('exported_at') or MISSING}",
        "",
        "压缩包里有：",
        f"  {DOCX_NAME}   核查初稿（Word）：文件信息、摘要、疑似片段表、全文时间轴转写、复核签字栏",
        f"  {CSV_NAME}        段落表格（UTF-8 带 BOM，可以用 Excel 直接打开）",
        f"  {JSON_NAME}       段落和处理信息（meta.notice 是上面的通知语）",
        clips_line,
        f"  {README_NAME}            本说明",
    ]
    return "\r\n".join(lines) + "\r\n"


def export_bundle(segments: list[dict], meta: dict, out_dir, clips_dir: str | None = None) -> str:
    """生成 Word 初稿、CSV、JSON、说明，连同 clips_dir 里的疑似片段打成 ZIP，返回 ZIP 的路径。

    - Word、CSV、JSON 也留在 out_dir 里（界面可以单独下载）；说明.txt 只放在 ZIP 里；
    - ZIP 的名字是 out_dir/<原文件名去扩展名>_review.zip（中文等字符换成 _）；
    - clips_dir 是步骤 7 export_clips 的输出文件夹，里面的 clip_*.wav 和 clips_index.csv 放进 ZIP 的 clips/；
      为 None（例如剧本文本演示，没有录音）时 ZIP 里留一个空的 clips/ 文件夹；
    - JSON 的 meta.notice 一定是通知语 NOTICE。传进来的段落和 meta 不会被修改。
    """
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    meta = dict(meta or {})
    meta["notice"] = NOTICE
    meta.setdefault("generated_by", GENERATED_BY)
    meta["exported_at"] = datetime.now().astimezone().isoformat(timespec="seconds")

    docx_path = out_dir / DOCX_NAME
    csv_path = out_dir / CSV_NAME
    json_path = out_dir / JSON_NAME
    build_docx(segments, meta, docx_path)
    write_csv(csv_path, segments)
    write_json(json_path, segments, meta)

    clip_files = []
    if clips_dir:
        clips_dir = Path(clips_dir)
        clip_files = sorted(clips_dir.glob("clip_*.wav"))
        if (clips_dir / CLIPS_INDEX_NAME).exists():
            clip_files.append(clips_dir / CLIPS_INDEX_NAME)
    clip_count = sum(1 for path in clip_files if path.suffix == ".wav")

    zip_path = out_dir / f"{_ascii_stem(meta.get('file'))}_review.zip"
    with zipfile.ZipFile(zip_path, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        zf.comment = GENERATED_BY.encode("utf-8")  # ZIP 的注释里也写明是机器生成的
        zf.write(docx_path, DOCX_NAME)
        zf.write(csv_path, CSV_NAME)
        zf.write(json_path, JSON_NAME)
        # utf-8-sig：带 BOM，老版本的 Windows 记事本也能认出是 UTF-8
        zf.writestr(README_NAME, _readme_text(meta, clip_count).encode("utf-8-sig"))
        if clip_files:
            for path in clip_files:
                zf.write(path, f"{CLIPS_FOLDER}/{path.name}")
        else:
            zf.writestr(f"{CLIPS_FOLDER}/", b"")  # 空文件夹
    return str(zip_path)
